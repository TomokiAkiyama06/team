"""抽出テキストと出典位置（ページ・スライド・見出し・行範囲）の対応を、
実際の抽出子プロセス → Indexer → Searcher → parse_answer の全経路で検証する。

各ページ/スライド/節に固有のマーカー文字列を入れ、
- マーカーを含むチャンクの location が正しいこと
- 1 つのチャンクに別ページのマーカーが混ざらないこと
- 検索結果から作った出典表示が正しいページを指すこと
を確かめる。
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from classbot.citations import parse_answer
from classbot.indexer import Indexer
from classbot.materials import MaterialStore
from classbot.search import Searcher


# --------------------------------------------------------------------------- 共通

def _index(db, settings, name: str, file_type: str) -> int:
    store = MaterialStore(db, settings)
    path = settings.materials_dir / name
    st = path.stat()
    mid = store.upsert_server(name, st.st_size, st.st_mtime, file_type)
    outcome = Indexer(db, settings, store).index_material(mid)
    assert outcome.status == "ready", outcome.message
    return mid


def _chunks(db, mid: int) -> list:
    return db.query(
        "SELECT c.* FROM chunks c JOIN materials m ON m.id = c.material_id "
        "AND c.version = m.active_version WHERE c.material_id = ? ORDER BY c.seq", (mid,))


def _locations_of(chunks, marker: str) -> set[str]:
    return {c["location"] for c in chunks if marker.lower() in c["text"].lower()}


def _cite(db, settings, query: str, marker: str) -> str:
    """marker を含む検索ヒットを [S?] として引用した回答の出典表示を返す。"""
    res = Searcher(db, settings).search(query)
    hit = next((h for h in res.hits if marker.lower() in h.text.lower()), None)
    assert hit is not None, f"{marker} が検索結果にない: {[h.text[:40] for h in res.hits]}"
    parsed = parse_answer(f"回答です[{hit.source_id}]\n[[SOURCES: {hit.source_id}]]", res.hits)
    assert parsed.invalid_ids == []
    return parsed.footer


# --------------------------------------------------------------------------- PDF

PDF_MARKERS = {1: "ALPHAPAGEONE", 2: "BRAVOPAGETWO", 4: "CHARLIEPAGEFOUR", 5: "DELTAPAGEFIVE"}


@pytest.fixture
def pdf_material(db, settings):
    import pymupdf

    doc = pymupdf.open()
    # p.1: 英語
    p = doc.new_page()
    p.insert_text((72, 72), f"Introduction {PDF_MARKERS[1]} course overview", fontsize=11)
    # p.2: 英語
    p = doc.new_page()
    p.insert_text((72, 72), f"Sorting algorithms {PDF_MARKERS[2]} quicksort", fontsize=11)
    # p.3: 空白ページ
    doc.new_page()
    # p.4: 日本語（CJK フォント）
    p = doc.new_page()
    p.insert_text((72, 72), f"第3回の課題はほうれん草の観察です {PDF_MARKERS[4]}", fontsize=11,
                  fontname="japan")
    # p.5: 長文（複数チャンクに分割される）
    p = doc.new_page()
    sentences = " ".join(f"{PDF_MARKERS[5]} sentence number {i} about recursion." for i in range(60))
    rc = p.insert_textbox(pymupdf.Rect(36, 36, 560, 800), sentences, fontsize=6)
    assert rc >= 0, "テキストがページに収まっていない"
    doc.save(settings.materials_dir / "第03回_講義資料.pdf")
    doc.close()
    return _index(db, settings, "第03回_講義資料.pdf", "pdf")


def test_pdf_each_marker_on_exact_page(db, settings, pdf_material):
    chunks = _chunks(db, pdf_material)
    for page, marker in PDF_MARKERS.items():
        assert _locations_of(chunks, marker) == {f"p.{page}"}, marker
    # 空白ページはチャンクを作らない
    assert not [c for c in chunks if c["location"] == "p.3"]
    # page 列も location と一致する
    for c in chunks:
        assert c["location"] == f"p.{c['page']}"


def test_pdf_no_chunk_mixes_pages(db, settings, pdf_material):
    for c in _chunks(db, pdf_material):
        present = [m for m in PDF_MARKERS.values() if m.lower() in c["text"].lower()]
        assert len(present) <= 1, (c["location"], present)


def test_pdf_long_page_splits_but_keeps_page(db, settings, pdf_material):
    p5 = [c for c in _chunks(db, pdf_material) if PDF_MARKERS[5].lower() in c["text"].lower()]
    assert len(p5) >= 2  # 800 文字を超えるので分割される
    assert {c["location"] for c in p5} == {"p.5"}
    assert all(len(c["text"]) <= settings.chunk_size_chars for c in p5)
    # 全文が欠けずにいずれかのチャンクに含まれる
    joined = " ".join(c["text"] for c in p5)
    for i in (0, 30, 59):
        assert f"sentence number {i} about" in joined


def test_pdf_japanese_page_extracted_and_cited(db, settings, pdf_material):
    chunks = _chunks(db, pdf_material)
    jp = [c for c in chunks if "ほうれん草" in c["text"]]
    assert jp and {c["location"] for c in jp} == {"p.4"}
    footer = _cite(db, settings, "ほうれん草の観察", "ほうれん草")
    assert footer.startswith("出典：第03回_講義資料.pdf p.4 [S")


def test_pdf_citation_points_to_right_page(db, settings, pdf_material):
    footer = _cite(db, settings, "quicksort sorting", PDF_MARKERS[2])
    assert re.fullmatch(r"出典：第03回_講義資料\.pdf p\.2 \[S\d+\]", footer), footer


# --------------------------------------------------------------------------- PPTX

def _make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    # スライド1: タイトル + 本文
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "授業ガイダンス"
    s.placeholders[1].text = "SLIDEONEBODY 評価方法"
    # スライド2: タイトル + 本文
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "第2回 ソート"
    s.placeholders[1].text = "SLIDETWOBODY クイックソート"
    # スライド3: 表
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "提出物一覧"
    table = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "課題"
    table.cell(0, 1).text = "形式"
    table.cell(1, 0).text = "TABLEMARKTHREE"
    table.cell(1, 1).text = "PDF"
    # スライド4: 本文なし・ノートのみ
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.notes_slide.notes_text_frame.text = "NOTEMARKFOUR 教員用メモ"
    prs.save(str(path))


@pytest.fixture
def pptx_path(settings):
    path = settings.materials_dir / "第02回_スライド.pptx"
    _make_pptx(path)
    return path


def test_pptx_slide_numbers_and_titles(db, settings, pptx_path):
    mid = _index(db, settings, pptx_path.name, "pptx")
    chunks = _chunks(db, mid)
    assert _locations_of(chunks, "SLIDEONEBODY") == {"スライド1"}
    assert _locations_of(chunks, "SLIDETWOBODY") == {"スライド2"}
    assert _locations_of(chunks, "TABLEMARKTHREE") == {"スライド3"}
    by_loc = {c["location"]: c for c in chunks}
    assert by_loc["スライド1"]["heading"] == "授業ガイダンス"
    assert by_loc["スライド3"]["heading"] == "提出物一覧"
    assert "TABLEMARKTHREE | PDF" in by_loc["スライド3"]["text"]
    # ノートは既定で含めない（教員用メモの漏えい防止）
    assert not _locations_of(chunks, "NOTEMARKFOUR")
    assert "スライド4" not in by_loc
    footer = _cite(db, settings, "クイックソート", "SLIDETWOBODY")
    assert footer.startswith("出典：第02回_スライド.pptx スライド2 [S")


def test_pptx_notes_enabled(db, settings, pptx_path):
    settings.pptx_include_notes = True
    mid = _index(db, settings, pptx_path.name, "pptx")
    assert _locations_of(_chunks(db, mid), "NOTEMARKFOUR") == {"スライド4"}


# --------------------------------------------------------------------------- DOCX

@pytest.fixture
def docx_material(db, settings):
    import docx
    from docx.enum.style import WD_STYLE_TYPE

    d = docx.Document()
    d.add_paragraph("PREAMBLEMARK この文書は授業の補足資料です。")
    d.add_paragraph("PREAMBLESECOND 二つ目の前文段落。")
    d.add_heading("第1章 概要", level=1)
    d.add_paragraph("CHAPTERONEMARK 授業の目的を説明します。")
    d.add_heading("課題の提出", level=2)
    d.add_paragraph("SUBMITMARK 提出は LMS から行います。")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "項目"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "TABLEDOCXMARK"
    table.cell(1, 1).text = "レポート"
    # 日本語名の見出しスタイル（日本語版 Word で作られた文書を想定）
    jp = d.styles.add_style("見出し 1", WD_STYLE_TYPE.PARAGRAPH)
    d.add_paragraph("第2章 評価", style=jp)
    d.add_paragraph("EVALMARK 期末試験とレポートで評価します。")
    d.save(str(settings.materials_dir / "シラバス.docx"))
    return _index(db, settings, "シラバス.docx", "docx")


def test_docx_heading_locations(db, settings, docx_material):
    chunks = _chunks(db, docx_material)
    assert _locations_of(chunks, "PREAMBLEMARK") == {"段落1付近"}
    assert _locations_of(chunks, "PREAMBLESECOND") == {"段落1付近"}
    assert _locations_of(chunks, "CHAPTERONEMARK") == {"見出し「第1章 概要」"}
    assert _locations_of(chunks, "SUBMITMARK") == {"見出し「課題の提出」"}
    # 表は直前の見出しの節に属する
    assert _locations_of(chunks, "TABLEDOCXMARK") == {"見出し「課題の提出」"}
    assert _locations_of(chunks, "EVALMARK") == {"見出し「第2章 評価」"}
    # 節をまたいだ混在がない
    for c in chunks:
        marks = [m for m in ("PREAMBLEMARK", "CHAPTERONEMARK", "SUBMITMARK", "EVALMARK") if m in c["text"]]
        assert len(marks) <= 1
    footer = _cite(db, settings, "期末試験 評価", "EVALMARK")
    assert footer.startswith("出典：シラバス.docx 見出し「第2章 評価」 [S")


# --------------------------------------------------------------------------- MD / TXT

def test_markdown_locations(db, settings):
    md = ("前置き MDPRE の行\n"
          "\n"
          "# 提出方法\n"
          "MDSUBMIT LMS から提出します。\n"
          "```\n"
          "# これはコード内のコメントで見出しではない MDCODE\n"
          "```\n"
          "## 締切\n"
          "MDDEADLINE 締切は翌週です。\n")
    (settings.materials_dir / "課題提出方法.md").write_text(md, encoding="utf-8")
    chunks = _chunks(db, _index(db, settings, "課題提出方法.md", "md"))
    assert _locations_of(chunks, "MDPRE") == {"L1-2"}
    assert _locations_of(chunks, "MDSUBMIT") == {"見出し「提出方法」"}
    assert _locations_of(chunks, "MDCODE") == {"見出し「提出方法」"}
    assert _locations_of(chunks, "MDDEADLINE") == {"見出し「締切」"}


def test_txt_line_ranges_contain_marker_line(db, settings):
    lines = []
    for i in range(1, 121):
        lines.append(f"TXTLINE{i:03d} " + "説明文" * 20)
        if i % 10 == 0:
            lines.append("")
    path = settings.materials_dir / "notes.txt"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    file_lines = path.read_text(encoding="utf-8").splitlines()
    chunks = _chunks(db, _index(db, settings, "notes.txt", "txt"))
    assert len({c["location"] for c in chunks}) > 1
    for c in chunks:
        m = re.fullmatch(r"L(\d+)-(\d+)", c["location"])
        assert m, c["location"]
        start, end = int(m.group(1)), int(m.group(2))
        for marker in re.findall(r"TXTLINE\d{3}", c["text"]):
            lineno = next(n for n, ln in enumerate(file_lines, 1) if ln.startswith(marker))
            assert start <= lineno <= end, (marker, lineno, c["location"])
