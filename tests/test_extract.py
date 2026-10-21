import json
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from classbot.extract import ExtractOptions, decode_text, extract_to_jsonl, iter_segments


def _read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]


def _make_pdf(path: Path, pages: list[str]) -> None:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((50, 72), text, fontname="japan", fontsize=11)
    doc.save(path)
    doc.close()


@pytest.fixture
def opts() -> ExtractOptions:
    return ExtractOptions(chunk_size_chars=400, chunk_overlap_chars=50)


def test_pdf_extract_pages(tmp_path: Path, settings):
    pdf = tmp_path / "第3回講義資料.pdf"
    _make_pdf(pdf, ["第3回 課題の説明です。レポートを提出してください。",
                    "提出期限は授業内で指示します。評価方法は出席と課題です。"])
    out = tmp_path / "out.jsonl"
    res = extract_to_jsonl(pdf, "pdf", settings, out)
    assert res.ok and res.status == "ready", res.error
    assert res.page_count == 2
    rows = _read_jsonl(out)
    assert {r["location"] for r in rows} == {"p.1", "p.2"}
    assert any("課題" in r["text"] and r["page"] == 1 for r in rows)
    assert res.chunk_count == len(rows)
    assert set(rows[0]) == {"seq", "text", "location", "page", "heading"}


def test_scanned_pdf_is_no_text(tmp_path: Path, settings):
    pdf = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 300, 300), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))
    doc.save(pdf)
    doc.close()
    out = tmp_path / "out.jsonl"
    res = extract_to_jsonl(pdf, "pdf", settings, out)
    assert not res.ok and res.status == "no_text"
    assert "テキストを抽出できませんでした" in res.error
    assert not out.exists()


def test_pdf_page_limit(tmp_path: Path, settings):
    settings.max_pages = 2
    pdf = tmp_path / "many.pdf"
    _make_pdf(pdf, ["本文です。"] * 3)
    res = extract_to_jsonl(pdf, "pdf", settings, tmp_path / "o.jsonl")
    assert res.status == "failed" and "上限" in res.error


def test_corrupt_pdf_does_not_crash_parent(tmp_path: Path, settings):
    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"%PDF-1.4\n" + b"\xff\x00garbage" * 100)
    res = extract_to_jsonl(pdf, "pdf", settings, tmp_path / "o.jsonl")
    assert not res.ok and res.status in ("failed", "no_text")


def test_timeout_kills_child(tmp_path: Path, settings):
    pdf = tmp_path / "a.pdf"
    _make_pdf(pdf, ["本文です。"])
    res = extract_to_jsonl(pdf, "pdf", settings, tmp_path / "o.jsonl", timeout=0.001)
    assert res.status == "failed" and "タイムアウト" in res.error


def test_pptx_slides_and_notes(tmp_path: Path, opts):
    path = tmp_path / "第02回_講義資料.pptx"
    prs = Presentation()
    for i in range(1, 4):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"タイトル{i}"
        slide.placeholders[1].text = f"スライド{i}の本文です。"
        slide.notes_slide.notes_text_frame.text = f"秘密のノート{i}"
    table_slide = prs.slides.add_slide(prs.slide_layouts[5])
    tbl = table_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "項目"
    tbl.cell(0, 1).text = "配点"
    prs.save(path)

    segs = list(iter_segments(path, "pptx", opts))
    assert [s.location for s in segs] == ["スライド1", "スライド2", "スライド3", "スライド4"]
    assert segs[0].heading == "タイトル1" and segs[0].page == 1
    assert "スライド2の本文" in segs[1].text
    assert all("秘密のノート" not in s.text for s in segs)
    assert "項目 | 配点" in segs[3].text

    opts.pptx_include_notes = True
    segs = list(iter_segments(path, "pptx", opts))
    assert "秘密のノート1" in segs[0].text


def test_pptx_full_pipeline(tmp_path: Path, settings):
    path = tmp_path / "s.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "演習の進め方"
    slide.placeholders[1].text = "各自でPythonの環境を用意してください。"
    prs.save(path)
    res = extract_to_jsonl(path, "pptx", settings, tmp_path / "o.jsonl")
    assert res.ok, res.error
    assert res.page_count == 1
    rows = _read_jsonl(res.chunks_path)
    assert rows[0]["location"] == "スライド1" and rows[0]["heading"] == "演習の進め方"


def test_docx_headings_and_tables(tmp_path: Path, opts):
    path = tmp_path / "シラバス.docx"
    d = Document()
    d.add_paragraph("前書きの段落です。")
    d.add_heading("成績評価", level=1)
    d.add_paragraph("期末レポート60%、小テスト40%で評価します。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "課題"
    t.cell(0, 1).text = "締切"
    t.cell(1, 0).text = "レポート1"
    t.cell(1, 1).text = "7月10日"
    d.add_heading("教科書", level=1)
    d.add_paragraph("指定なし。")
    d.save(path)

    segs = list(iter_segments(path, "docx", opts))
    locs = [s.location for s in segs]
    assert locs[0] == "段落1付近"
    assert "見出し「成績評価」" in locs and "見出し「教科書」" in locs
    grading = next(s for s in segs if s.heading == "成績評価")
    assert "期末レポート60%" in grading.text
    assert "レポート1 | 7月10日" in grading.text


def test_markdown_headings(tmp_path: Path, opts):
    path = tmp_path / "課題提出方法.md"
    path.write_text("最初の説明\n\n# 提出方法\nLMSから提出します。\n\n```\n# コード内のコメント\n```\n"
                    "## 締切\n毎週金曜日です。\n", encoding="utf-8")
    segs = list(iter_segments(path, "md", opts))
    assert [s.location for s in segs] == ["L1-2", "見出し「提出方法」", "見出し「締切」"]
    assert "# コード内のコメント" in segs[1].text


def test_txt_line_ranges_and_cp932(tmp_path: Path, settings):
    path = tmp_path / "連絡.txt"
    body = "\n\n".join(f"段落{i}：" + "連絡事項です。" * 20 for i in range(10))
    path.write_bytes(body.encode("cp932"))
    res = extract_to_jsonl(path, "txt", settings, tmp_path / "o.jsonl")
    assert res.ok, res.error
    rows = _read_jsonl(res.chunks_path)
    assert all(r["location"].startswith("L") for r in rows)
    assert "連絡事項" in rows[0]["text"]
    assert rows[0]["location"].startswith("L1-")


def test_decode_text():
    assert decode_text("日本語".encode("utf-8-sig")) == "日本語"
    assert decode_text("日本語".encode("cp932")) == "日本語"


def test_max_extracted_chars_truncates(tmp_path: Path, settings):
    settings.max_extracted_chars = 1000
    path = tmp_path / "long.txt"
    path.write_text("\n\n".join("長い文章です。" * 30 for _ in range(50)), encoding="utf-8")
    res = extract_to_jsonl(path, "txt", settings, tmp_path / "o.jsonl")
    assert res.ok and res.error and "上限" in res.error
    total = sum(len(r["text"]) for r in _read_jsonl(res.chunks_path))
    assert total <= 1000


def test_empty_text_file_no_text(tmp_path: Path, settings):
    path = tmp_path / "blank.md"
    path.write_text("\n\n   \n", encoding="utf-8")
    res = extract_to_jsonl(path, "md", settings, tmp_path / "o.jsonl")
    assert res.status == "no_text"
