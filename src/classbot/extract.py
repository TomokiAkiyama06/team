"""資料からのテキスト抽出。

抽出は子プロセス（spawn）で行い、結果のチャンクを JSONL ファイルへ書き出す。
- 不正なファイルでライブラリがクラッシュしてもボット本体は落ちない
- タイムアウト時は子プロセスを kill する
- RLIMIT_AS でメモリ使用量に上限を設ける
- チャンクはファイルへ逐次書き出すため、長大な資料でも親プロセスのメモリは増えない
"""

from __future__ import annotations

import json
import logging
import multiprocessing
import re
import traceback
from collections.abc import Iterator
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from classbot.chunking import Segment, chunk_segments

log = logging.getLogger(__name__)

# スキャン PDF とみなす閾値（資料全体の非空白文字数）
_NO_TEXT_MIN_CHARS = 10
# OCR を試みるページの閾値（ページ内の非空白文字数）
_OCR_PAGE_MIN_CHARS = 20
# 見出しの無い長い節を分割するセグメント長（文字）
_SEGMENT_FLUSH_CHARS = 20000

_MD_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_FENCE_RE = re.compile(r"^\s{0,3}(```|~~~)")


@dataclass
class ExtractOptions:
    """子プロセスへ渡す抽出設定（秘密情報を含む Settings 全体は渡さない）。"""

    chunk_size_chars: int = 800
    chunk_overlap_chars: int = 120
    max_pages: int = 2000
    max_extracted_chars: int = 5_000_000
    pptx_include_notes: bool = False
    ocr_enabled: bool = False
    ocr_language: str = "jpn+eng"
    extract_max_memory_mb: int = 1536

    @classmethod
    def from_settings(cls, settings: Any) -> ExtractOptions:
        return cls(**{k: getattr(settings, k) for k in cls.__dataclass_fields__
                      if hasattr(settings, k)})


@dataclass
class ExtractResult:
    ok: bool
    status: str  # "ready" | "no_text" | "failed"
    chunks_path: Path | None
    page_count: int | None
    chunk_count: int
    error: str | None


class ExtractError(Exception):
    """利用者に表示できる日本語メッセージを持つ抽出エラー。"""


# ---------------------------------------------------------------------------
# テキストのデコード
# ---------------------------------------------------------------------------

def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp932"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ---------------------------------------------------------------------------
# 形式別のセグメント抽出
# ---------------------------------------------------------------------------

def _iter_pdf(path: Path, opts: Any) -> Iterator[Segment]:
    import pymupdf

    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise ExtractError("パスワード保護された PDF は処理できません")
        if doc.page_count > opts.max_pages:
            raise ExtractError(f"ページ数が上限（{opts.max_pages}ページ）を超えています")
        for i in range(doc.page_count):
            page = doc.load_page(i)
            text = page.get_text("text")
            if opts.ocr_enabled and len("".join(text.split())) < _OCR_PAGE_MIN_CHARS:
                text = _ocr_page(page, opts) or text
            yield Segment(text=text, location=f"p.{i + 1}", page=i + 1)


def _ocr_page(page: Any, opts: Any) -> str:
    """ローカルの Tesseract で OCR する（外部 API へ画像を送らない）。"""
    try:
        tp = page.get_textpage_ocr(language=opts.ocr_language, dpi=200, full=True)
    except Exception as e:  # Tesseract 未導入・言語データ不足など
        raise ExtractError(
            "OCR を実行できませんでした（Tesseract と言語データがインストールされているか確認してください）"
        ) from e
    return page.get_text("text", textpage=tp)


def _pptx_shape_texts(shapes: Any) -> Iterator[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_shape_texts(shape.shapes)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            t = shape.text_frame.text
            if t.strip():
                yield t
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    yield " | ".join(cells)


def _iter_pptx(path: Path, opts: Any) -> Iterator[Segment]:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = list(prs.slides)
    if len(slides) > opts.max_pages:
        raise ExtractError(f"スライド数が上限（{opts.max_pages}枚）を超えています")
    for i, slide in enumerate(slides, start=1):
        title = None
        try:
            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except Exception:
            title = None
        parts = list(_pptx_shape_texts(slide.shapes))
        if opts.pptx_include_notes and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            if notes.strip():
                parts.append("［ノート］" + notes)
        yield Segment(text="\n".join(parts), location=f"スライド{i}", page=i, heading=title)


def _docx_heading_level(style_name: str) -> bool:
    s = (style_name or "").lower()
    return s.startswith("heading") or s.startswith("見出し") or s == "title" or s == "表題"


def _iter_docx(path: Path, opts: Any) -> Iterator[Segment]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    body = document.element.body
    heading: str | None = None
    buf: list[str] = []
    buf_len = 0
    para_no = 0
    start_para = 1

    def flush() -> Segment | None:
        nonlocal buf, buf_len
        if not buf:
            return None
        loc = f"見出し「{heading}」" if heading else f"段落{start_para}付近"
        seg = Segment(text="\n".join(buf), location=loc, heading=heading)
        buf, buf_len = [], 0
        return seg

    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para_no += 1
            p = Paragraph(child, document)
            text = p.text
            style = p.style.name if p.style is not None else ""
            if _docx_heading_level(style) and text.strip():
                seg = flush()
                if seg:
                    yield seg
                heading = text.strip()[:80]
                start_para = para_no
                buf.append(text)
                buf_len += len(text)
                continue
            if not buf:
                start_para = para_no
            if text.strip():
                buf.append(text)
                buf_len += len(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = []
                for c in row.cells:
                    t = c.text.strip()
                    if t and (not cells or cells[-1] != t):  # 結合セルの重複を除く
                        cells.append(t)
                if cells:
                    line = " | ".join(cells)
                    buf.append(line)
                    buf_len += len(line)
        if buf_len >= _SEGMENT_FLUSH_CHARS:
            seg = flush()
            if seg:
                yield seg
            start_para = para_no + 1
    seg = flush()
    if seg:
        yield seg


def _iter_markdown(text: str) -> Iterator[Segment]:
    heading: str | None = None
    buf: list[str] = []
    start = 1
    in_fence = False
    for lineno, line in enumerate(text.splitlines(), start=1):
        if _FENCE_RE.match(line):
            in_fence = not in_fence
        m = None if in_fence else _MD_HEADING_RE.match(line)
        if m:
            if any(s.strip() for s in buf):
                yield _md_segment(buf, heading, start, lineno - 1)
            heading = m.group(2).strip()[:80]
            buf = [heading]
            start = lineno
            continue
        buf.append(line)
        if sum(len(s) for s in buf) >= _SEGMENT_FLUSH_CHARS:
            yield _md_segment(buf, heading, start, lineno)
            buf, start = [], lineno + 1
    if any(s.strip() for s in buf):
        yield _md_segment(buf, heading, start, start + len(buf) - 1)


def _md_segment(buf: list[str], heading: str | None, start: int, end: int) -> Segment:
    loc = f"見出し「{heading}」" if heading else f"L{start}-{end}"
    return Segment(text="\n".join(buf), location=loc, heading=heading)


def _iter_txt(text: str, size: int) -> Iterator[Segment]:
    """空行区切りの段落を size 文字程度にまとめ、行範囲を位置情報にする。"""
    buf: list[str] = []
    buf_len = 0
    start = 1
    lines = text.splitlines()
    for lineno, line in enumerate(lines, start=1):
        if not buf and not line.strip():
            start = lineno + 1
            continue
        buf.append(line)
        buf_len += len(line)
        if (not line.strip() and buf_len >= size) or buf_len >= size * 2:
            yield Segment(text="\n".join(buf), location=f"L{start}-{lineno}")
            buf, buf_len, start = [], 0, lineno + 1
    if any(s.strip() for s in buf):
        yield Segment(text="\n".join(buf), location=f"L{start}-{start + len(buf) - 1}")


def iter_segments(path: Path, file_type: str, settings: Any) -> Iterator[Segment]:
    """資料をセグメント（ページ・スライド・見出し節・行範囲）単位で抽出する。"""
    if file_type == "pdf":
        yield from _iter_pdf(path, settings)
    elif file_type == "pptx":
        yield from _iter_pptx(path, settings)
    elif file_type == "docx":
        yield from _iter_docx(path, settings)
    elif file_type in ("txt", "md"):
        text = decode_text(path.read_bytes())
        if file_type == "md":
            yield from _iter_markdown(text)
        else:
            yield from _iter_txt(text, settings.chunk_size_chars)
    else:
        raise ExtractError("未対応の形式です")


def _page_count(path: Path, file_type: str) -> int | None:
    try:
        if file_type == "pdf":
            import pymupdf

            with pymupdf.open(path) as doc:
                return doc.page_count
        if file_type == "pptx":
            from pptx import Presentation

            return len(Presentation(str(path)).slides)
    except Exception:
        return None
    return None


# ---------------------------------------------------------------------------
# 子プロセス
# ---------------------------------------------------------------------------

def _apply_memory_limit(mb: int) -> None:
    try:
        import resource

        limit = mb * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (limit, limit))
    except (ImportError, ValueError, OSError):
        pass


def _run_extraction(path: Path, file_type: str, opts: ExtractOptions, out_path: Path) -> dict:
    total_chars = 0
    nonspace_chars = 0
    count = 0
    truncated = False
    with open(out_path, "w", encoding="utf-8") as out:
        segments = iter_segments(path, file_type, opts)
        for chunk in chunk_segments(segments, opts.chunk_size_chars, opts.chunk_overlap_chars):
            if total_chars + len(chunk.text) > opts.max_extracted_chars:
                truncated = True
                break
            total_chars += len(chunk.text)
            nonspace_chars += len("".join(chunk.text.split()))
            out.write(json.dumps(asdict(chunk), ensure_ascii=False) + "\n")
            count += 1
    page_count = _page_count(path, file_type)
    if nonspace_chars < _NO_TEXT_MIN_CHARS:
        if file_type == "pdf":
            msg = ("テキストを抽出できませんでした（スキャンされた画像のみの PDF の可能性があります）"
                   if not opts.ocr_enabled else "OCR を実行しましたがテキストを抽出できませんでした")
        else:
            msg = "テキストを抽出できませんでした（本文が含まれていません）"
        return {"status": "no_text", "page_count": page_count, "chunk_count": 0, "error": msg}
    err = (f"抽出文字数が上限（{opts.max_extracted_chars}文字）に達したため、以降を省略しました"
           if truncated else None)
    return {"status": "ready", "page_count": page_count, "chunk_count": count, "error": err}


def _preload_libraries(file_type: str) -> None:
    """メモリ制限を掛ける前にライブラリを読み込んでおく。

    PyMuPDF 2.x は任意依存の pymupdf.layout（onnxruntime）を自動 import し、
    スレッドプール用に大きな仮想メモリを確保するため RLIMIT_AS 下で停止することがある。
    テキスト抽出には不要なので import を無効化する。
    """
    import sys

    sys.modules.setdefault("pymupdf.layout", None)  # type: ignore[arg-type]
    if file_type == "pdf":
        import pymupdf  # noqa: F401
    elif file_type == "pptx":
        import pptx  # noqa: F401
    elif file_type == "docx":
        import docx  # noqa: F401


def _child_main(path: str, file_type: str, opts_dict: dict, out_path: str, result_path: str) -> None:
    opts = ExtractOptions(**opts_dict)
    _preload_libraries(file_type)
    _apply_memory_limit(opts.extract_max_memory_mb)
    try:
        result = _run_extraction(Path(path), file_type, opts, Path(out_path))
    except ExtractError as e:
        result = {"status": "failed", "page_count": None, "chunk_count": 0, "error": str(e)}
    except MemoryError:
        result = {"status": "failed", "page_count": None, "chunk_count": 0,
                  "error": "メモリ上限を超えたため処理を中止しました"}
    except Exception as e:  # 壊れたファイル等
        result = {"status": "failed", "page_count": None, "chunk_count": 0,
                  "error": f"ファイルを解析できませんでした（{type(e).__name__}）",
                  "detail": traceback.format_exc(limit=5)}
    with open(result_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False)


def extract_to_jsonl(path: Path, file_type: str, settings: Any, out_path: Path,
                     timeout: float | None = None) -> ExtractResult:
    """子プロセスで抽出・チャンク化し、out_path に JSONL（1行1チャンク）を書き出す。

    JSONL の各行: {"seq", "text", "location", "page", "heading"}
    """
    opts = settings if isinstance(settings, ExtractOptions) else ExtractOptions.from_settings(settings)
    if timeout is None:
        timeout = float(getattr(settings, "extract_timeout_sec", 300))
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    result_path = out_path.with_name(out_path.name + ".result.json")
    for p in (out_path, result_path):
        p.unlink(missing_ok=True)

    def _fail(msg: str) -> ExtractResult:
        out_path.unlink(missing_ok=True)
        result_path.unlink(missing_ok=True)
        return ExtractResult(False, "failed", None, None, 0, msg)

    ctx = multiprocessing.get_context("spawn")
    proc = ctx.Process(target=_child_main,
                       args=(str(path), file_type, asdict(opts), str(out_path), str(result_path)),
                       daemon=True)
    proc.start()
    proc.join(timeout)
    if proc.is_alive():
        proc.kill()
        proc.join(10)
        log.warning("extraction timed out", extra={"file_type": file_type, "timeout": timeout})
        return _fail(f"処理がタイムアウトしました（{int(timeout)}秒）")
    if not result_path.exists():
        log.warning("extraction process crashed", extra={"exitcode": proc.exitcode})
        return _fail("抽出処理が異常終了しました（ファイルが壊れているか、メモリ上限を超えた可能性があります）")
    try:
        data = json.loads(result_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return _fail("抽出結果を読み込めませんでした")
    result_path.unlink(missing_ok=True)
    if data.get("detail"):
        log.warning("extraction failed", extra={"detail": data["detail"]})
    status = data["status"]
    if status != "ready":
        out_path.unlink(missing_ok=True)
        return ExtractResult(False, status, None, data.get("page_count"), 0, data.get("error"))
    return ExtractResult(True, "ready", out_path, data.get("page_count"), data["chunk_count"],
                         data.get("error"))
